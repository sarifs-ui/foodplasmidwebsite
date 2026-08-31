import openpyxl
import json

# Check All_Data.xlsx
wb = openpyxl.load_workbook("All_Data.xlsx", data_only=True, read_only=True)
ws = wb.active
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print("Headers:", headers[:20])

kefir_rows = []
boza_rows = []
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    row_dict = dict(zip(headers, row))
    for key, val in row_dict.items():
        if val and ("kefir" in str(val).lower() or "boza" in str(val).lower()):
            if "kefir" in str(val).lower():
                kefir_rows.append((key, val, {k: v for k, v in row_dict.items() if v}))
            if "boza" in str(val).lower():
                boza_rows.append((key, val, {k: v for k, v in row_dict.items() if v}))
    if i > 10000:
        break

print(f"\nKefir rows ({len(kefir_rows)}):")
for col, val, row in kefir_rows[:5]:
    print(f"  Column '{col}' = '{val}': {dict(list(row.items())[:8])}")

print(f"\nBoza rows ({len(boza_rows)}):")
for col, val, row in boza_rows[:5]:
    print(f"  Column '{col}' = '{val}': {dict(list(row.items())[:8])}")

wb.close()

# Check gfpr.json for kefir/boza
with open("gfpr.json", "r") as f:
    data = json.load(f)

print(f"\nTotal samples in gfpr.json: {len(data)}")
kefir_in_json = [d for d in data if any("kefir" in str(v).lower() for v in d.values())]
boza_in_json = [d for d in data if any("boza" in str(v).lower() for v in d.values())]
print(f"Kefir in gfpr.json: {len(kefir_in_json)}")
print(f"Boza in gfpr.json: {len(boza_in_json)}")

# Check cazyme ID format vs gfpr.json Run_ID
with open("cazyme.tsv", "r", encoding="utf-8", errors="replace") as f:
    content = f.read()
lines = content.strip().split('\n')
header = lines[0].lstrip('\ufeff').split(';')
cazyme_ids = set()
for line in lines[1:]:
    parts = line.split(';')
    if parts:
        cazyme_ids.add(parts[0])

all_run_ids = set(d['run_id'] for d in data if d.get('run_id'))
all_sample_ids = set(d['sample_id'] for d in data if d.get('sample_id'))
print(f"\nTotal cazyme unique IDs: {len(cazyme_ids)}")
print(f"Overlap with run_id: {len(cazyme_ids & all_run_ids)}")
print(f"Overlap with sample_id: {len(cazyme_ids & all_sample_ids)}")
print(f"Sample cazyme IDs: {list(cazyme_ids)[:5]}")

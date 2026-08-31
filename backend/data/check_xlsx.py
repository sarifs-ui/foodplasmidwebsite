import openpyxl
import json

wb = openpyxl.load_workbook("All_Data.xlsx", data_only=True, read_only=True)
ws = wb.active
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print("Headers:", headers[:20])

total = 0
no_sample_id = []
in_house_rows = []
null_sample_id = []

for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
    if not any(v is not None for v in row):
        continue  # skip completely empty rows
    row_dict = dict(zip(headers, row))
    total += 1
    sid = row_dict.get("Sample_ID")
    rid = row_dict.get("Run_ID")
    
    if not sid:
        null_sample_id.append({"row": i+2, "Run_ID": rid, "Category": row_dict.get("Category"), "Type": row_dict.get("Type")})
    
    if rid and "IN_HOUSE_ID_" in str(rid):
        in_house_rows.append({"Sample_ID": sid, "Run_ID": rid, "Category": row_dict.get("Category"), "Type": row_dict.get("Type")})

print(f"\nTotal non-empty rows in xlsx: {total}")
print(f"Rows with null Sample_ID: {len(null_sample_id)}")
print(f"Rows with IN_HOUSE_ID_ in Run_ID: {len(in_house_rows)}")

print(f"\nFirst 10 null Sample_ID rows:")
for r in null_sample_id[:10]:
    print(f"  Row {r['row']}: Run_ID={r['Run_ID']}, Category={r['Category']}, Type={r['Type']}")

print(f"\nIN_HOUSE_ID_ rows (first 10):")
for r in in_house_rows[:10]:
    print(f"  Sample_ID={r['Sample_ID']}, Run_ID={r['Run_ID']}, Category={r['Category']}, Type={r['Type']}")

wb.close()

# Compare with gfpr.json
with open("gfpr.json", "r") as f:
    data = json.load(f)
print(f"\ngfpr.json total: {len(data)}")
in_house_in_json = [d for d in data if d.get("run_id", "").startswith("IN_HOUSE_ID_") if d.get("run_id")]
print(f"IN_HOUSE_ID_ in gfpr.json: {len(in_house_in_json)}")
